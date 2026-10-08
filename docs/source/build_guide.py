from pathlib import Path
import json, shutil
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
B=Path(__file__).parent;R=B.parent.parent; A=B/'generated';A.mkdir(exist_ok=True)
pdfmetrics.registerFont(TTFont('Chinese',r'C:\Windows\Fonts\msjh.ttc'))
W,H=1200,800
TEAL='#137F83';INK='#183A3B';RED='#C92B46';PALE='#F0F6F3'
def illustration(name,header,draw):
 c=canvas.Canvas(str(A/(name+'.pdf')),pagesize=(W,H));c.setFillColorRGB(1,1,1);c.rect(0,0,W,H,fill=1,stroke=0)
 def text(x,y,s,size=32,color=INK):
  c.setFillColor(color);c.setFont('Chinese',size);c.drawString(x,H-y,s)
 def box(x,y,w,h,fill=PALE,stroke=None):
  c.setFillColor(fill);c.setStrokeColor(stroke or fill);c.setLineWidth(4 if stroke else 1);c.roundRect(x,H-y-h,w,h,18,fill=1,stroke=1)
 def call(x,y,s): text(x,y,s,30,RED)
 text(45,58,header,36);text(940,58,'操作示意',25,'#677B7B');draw(c,text,box,call);c.save()
 doc=pdfium.PdfDocument(str(A/(name+'.pdf')));doc[0].render(scale=1.4).to_pil().save(A/(name+'.png'));doc.close()
 return A/(name+'.png')

def start(c,t,b,k):
 b(55,110,1090,235);t(90,160,'Agent 會先問你',28,TEAL);t(90,215,'你有 Google、ChatGPT、Cloudflare 帳號嗎？',36)
 t(90,285,'你：有 Google 和 ChatGPT，還沒有 Cloudflare。',30)
 b(55,395,500,280);t(90,453,'交給 Agent',36);t(90,515,'下載的完整設定包',34);t(90,580,'請讀 AGENTS.md，帶我設定。',28)
 b(610,395,535,280);t(645,453,'你接手時',36);t(645,515,'登入 → 授權 → 確認結果',30);t(645,580,'每次只做一個動作',32)
 k(65,742,'先下載 ZIP，解壓縮，再用 Agent 開啟整個資料夾。')
def signup(c,t,b,k):
 b(185,102,830,585,'#FFFFFF','#CFDCD8');t(225,158,'Create your Cloudflare account',36)
 t(230,226,'Email',30);b(220,245,750,68,'#FFFFFF',RED);t(245,290,'填你的電子郵件地址',30);k(980,290,'1')
 t(230,359,'Password',30);b(220,380,750,68,'#FFFFFF',RED);t(245,425,'由你自己設定密碼',30);k(980,425,'2')
 b(220,503,750,70,TEAL,RED);t(455,551,'Create Account',34,'#FFFFFF');k(980,550,'3')
 t(230,632,'註冊後，到剛才的信箱按驗證連結。',30)
 t(65,743,'免費註冊；不必先買網域或選付費方案。',32,TEAL)
def google(c,t,b,k):
 for x,title,lines in [(50,'專用寄件 Gmail',['只用來寄每天的提醒','Google 會要求完整 Gmail 權限','用另一個帳號，保護平常的信箱']), (625,'主帳號',['平常收信、放行事曆的帳號','授權日曆事件讀寫','每筆新增或改期仍須你確認'])]:
  b(x,135,525,420);t(x+30,200,title,38,TEAL)
  for i,s in enumerate(lines):t(x+30,285+i*82,s,27)
 k(80,635,'看到 Google 選帳號畫面時，先確認這次要用哪一個。')
 t(80,705,'Agent 會先準備好 Google Cloud 設定並開啟授權頁。',30)
def project(c,t,b,k):
 b(165,110,870,560,'#FFFFFF','#CFDCD8');t(205,175,'建立專案',40)
 t(205,245,'專案名稱',30);b(195,266,795,72,'#FFFFFF',RED);t(225,314,'AI 秘書',34)
 t(205,398,'記憶',30);b(195,420,795,70,'#FFFFFF',RED);t(225,468,'預設記憶',34)
 t(205,563,'專案指示：貼上 Agent 為你填好的版本',30)
 k(165,740,'只新增這個專案；不用修改你原來的其他專案。')
def mcp(c,t,b,k):
 b(155,110,890,600,'#FFFFFF','#CFDCD8');t(195,170,'建立自訂 MCP 伺服器',38)
 t(195,240,'名稱',27);b(185,260,790,65,'#FFFFFF',RED);t(210,305,'AI 秘書進度',32)
 t(195,380,'連線  伺服器 URL',27);b(185,400,790,70,'#FFFFFF',RED);t(210,447,'貼上 Agent 提供的你的服務網址 /mcp',28)
 t(195,525,'身分驗證',27);b(185,545,790,65,'#FFFFFF',RED);t(210,590,'OAuth',32)
 t(195,674,'確認你信任這個服務，再按建立與連接。',28)
def approval(c,t,b,k):
 b(135,115,930,565,'#FFFFFF','#CFDCD8');t(175,180,'連結 AI 秘書進度',40)
 t(175,254,'把頁面下方「請求識別」交給 Agent。',32)
 t(175,330,'Agent 會把一次性連結碼放到你的剪貼簿。',30)
 b(165,375,825,76,'#FFFFFF',RED);t(193,426,'你在這裡貼上連結碼',32)
 b(165,492,825,76,TEAL,RED);t(430,543,'授權這次連線',34,'#FFFFFF')
 t(175,640,'連結碼五分鐘有效；不要貼進聊天或截圖。',28)
def mail(c,t,b,k):
 b(80,112,1040,550,'#FFFFFF','#CFDCD8');t(120,175,'AI 秘書  總覽驗收',40)
 t(120,245,'急件  近期截止  一般待辦  今日行程',30,TEAL)
 b(112,290,960,195,'#FFFFFF',RED);t(145,352,'□  【教材示範】寄出課程報名資料',34)
 t(170,430,'勾選後 → 等到「完成已保存」',32,TEAL)
 t(120,565,'關閉，再重開同一封信，確認完成狀態還在。',30)
 k(90,740,'Agent 會讀回進度並清理示範，你不用操作資料庫。')
def calendar(c,t,b,k):
 b(50,110,1100,135);t(80,160,'你：明天下午幫我留一小時準備資料。',32);t(80,210,'Agent：先查日曆並提出一個時間。',30)
 b(100,290,1000,225,'#FFFFFF',RED);t(140,345,'安排提案',36,TEAL)
 t(140,405,'準備報名資料  ·  10 月 20 日  14:00–15:00',32)
 t(140,470,'這個時段沒有撞期。要幫你排進去嗎？',32)
 t(110,590,'你：可以，請排入。',34);t(110,664,'Agent：寫入後讀回，確認是同一筆行程。',30)
 t(110,738,'日期為示例，請換成你真的想安排的一件事。',27,'#677B7B')
def daily(c,t,b,k):
 b(90,115,1020,285);t(145,185,'設定與驗收完成',42);t(145,255,'信件勾選會保存  ·  行程新增與改期已核對',30)
 t(145,332,'Agent 已清理教材示範，等待你確認啟用。',30)
 b(90,455,1020,190,'#FFFFFF',RED);t(145,525,'你：驗收沒問題，請啟用每天的提醒。',34);t(145,600,'每天 09:00  ·  台灣時間  ·  一封 Gmail',34,TEAL)
 t(95,742,'下一次早上 9 點後，實際收信才算首次排程驗收。',29)
def use(c,t,b,k):
 samples=[('交辦','週五前寄出資料，幫我記起來。'),('安排','找一個小時準備，先檢查撞期。'),('更新','我做完了，接下來追蹤對方回覆。')]
 for i,(a,z) in enumerate(samples):
  y=125+i*172;b(55,y,1090,138);t(90,y+53,a,31,TEAL);t(90,y+108,z,36)
 k(70,730,'想停每天的信，跟 Agent 說「請暫停每日提醒」。')
figs={n:illustration(n,h,f) for n,h,f in [('01-agent','把設定包交給 Agent',start),('02-cloudflare','Cloudflare 免費註冊',signup),('04-google','Google 授權先確認帳號',google),('05-project','建立 ChatGPT 專案',project),('06-mcp','連上自己的秘書服務',mcp),('07-approval','本人完成連線授權',approval),('09-mail','用一封信確認完成會保存',mail),('10-calendar','日曆安排先確認再寫入',calendar),('11-daily','開啟每天早上九點的提醒',daily),('12-use','平常這樣交辦',use)]}
shots=B/'screenshots'
pages=[
('第 1 步 把設定包交給 Agent','目的：讓 Agent 問帳號、做設定，你只在需要時接手。',figs['01-agent'],None,'下載並解壓 ZIP，用 Codex 或 Claude Code 開啟整個資料夾，再貼：「請讀 AGENTS.md，帶我設定自己的 AI 秘書。」','目前支援 Windows。先核對 ChatGPT 有自訂連接功能，再開始雲端設定。'),
('第 2 步 申請 Cloudflare','目的：讓提醒每天在雲端執行，電腦關機也能寄信。',figs['02-cloudflare'],None,'Agent 開啟 dash.cloudflare.com/sign-up。照紅框填電子郵件、設定密碼，按 Create Account，再到信箱驗證。','已有 Cloudflare 就跳過。這是依官方步驟繪製的示意圖，實際畫面可能略有不同。'),
('第 3 步 授權部署工具','目的：讓 Agent 在你的 Cloudflare 建立秘書服務。',shots/'14-wrangler-permissions.png',(375,105,960,525),'確認頁面是 Cloudflare 官方的 Wrangler、右側是你的帳號，閱讀權限後按 Authorize（授權）。','Agent 接著建立服務和資料庫。使用免費方案，不需要先買網域。'),
('第 4 步 完成 Google 授權','目的：讓秘書寄提醒，並管理你指定的行事曆。',figs['04-google'],None,'跟著 Agent 分兩次選帳號：寄信選專用 Gmail；日曆選平常的主帳號。看完各次權限，再由你按同意。','寄件 Gmail 需另備一個。SMTP 要求完整 Gmail 權限；主帳號只做日曆授權。'),
('第 5 步 建立 AI 秘書專案','目的：讓 ChatGPT 知道怎麼當你的秘書。',figs['05-project'],None,'在 ChatGPT 新增「AI 秘書」專案，記憶選「預設記憶」。打開專案指示，貼上 Agent 準備好的內容。','這是操作示意。若你看不到這些功能，先回報 Agent 核對帳號功能。'),
('第 6 步 加入秘書連線','目的：讓 ChatGPT 能保存待辦、讀寫你的日曆。',figs['06-mcp'],None,'外掛程式 → 新增 → 建立自訂 MCP 伺服器。按圖填名稱、Agent 給你的完整 /mcp 網址，身分驗證選 OAuth。','這是操作示意；不要照抄別人的服務網址。登入與確認由本人完成。'),
('第 7 步 確認這次連線','目的：只讓你本人授權 ChatGPT 使用秘書服務。',figs['07-approval'],None,'把「請求識別」交給 Agent。等它準備好後，你把剪貼簿中的連結碼貼入欄位，按「授權這次連線」。','接著在工作模式選「AI 秘書」專案及「AI 秘書進度」外掛。Agent 會測試連線。'),
('第 8 步 開啟信內勾選','目的：在 Gmail 裡勾一下，就能保存完成進度。',shots/'23-gmail-dynamic-sender.png',(360,186,565,350),'主帳號 Gmail：設定 → 查看所有設定 → 一般 → 動態電子郵件 → 開發人員設定。照紅框填專用寄件 Gmail，再儲存。','先勾選「啟用動態電子郵件」。使用 Gmail 網頁或 App；iPhone 內建郵件不能信內勾選。'),
('第 9 步 收一封驗收信','目的：確認信收得到，完成狀態也真的有保存。',figs['09-mail'],None,'Agent 建好示範待辦後，確認只寄一封驗收信。你在 Gmail 勾選，等「完成已保存」，關閉再重開確認。','Agent 會核對聊天讀回的進度，確認下一封不再列出已完成項目，再清理示範。'),
('第 10 步 確認日曆安排','目的：確認先查撞期，再經你同意排進日曆。',figs['10-calendar'],None,'請秘書安排一件你真的要做的事。核對日期、時間與撞期結果，確認後讓它寫入，再到 Google Calendar 看同一筆。','再指定一次改期並確認，核對仍是同一筆。多人會議、整組重複規則或刪除，直接到日曆操作。'),
('第 11 步 開啟每天提醒','目的：從明天開始，每天早上收到一封總覽。',figs['11-daily'],None,'確認信件和日曆都驗收完成，再跟 Agent 說：「請啟用每天早上九點的提醒。」','第一次到點後，實際核對收件匣；没收到先看垃圾郵件，再請 Agent 查原因。'),
('第 12 步 平常交辦就好','目的：少花時間整理，把注意力留給要做的事。',figs['12-use'],None,'在「AI 秘書」專案裡，一件活動開一個聊天。告訴它下一步和期限；做好、延期或取消，也在聊天裡說。','需要排進日曆時，等提案後再確認。每天 Gmail 會整理急件、近期截止、一般待辦和今日行程。')]
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
for name in ['Normal','Title','Heading 1','Caption']:
 s=doc.styles[name];s.font.name='Microsoft JhengHei';s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft JhengHei');s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(12)
doc.styles['Normal'].font.size=Pt(13);doc.styles['Normal'].paragraph_format.line_spacing=1.28
doc.styles['Heading 1'].font.size=Pt(23);doc.styles['Title'].font.size=Pt(23);doc.styles['Caption'].font.size=Pt(10)
head=sec.header.paragraphs[0];head.text='AI 秘書操作圖解';head.style='Caption'
foot=sec.footer.paragraphs[0];foot.text='AI 秘書  ·  2026 年 10 月 8 日                                     ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for i,(title,purpose,img,crop,action,note) in enumerate(pages):
 if i:doc.add_page_break()
 doc.add_paragraph(title,'Title' if i==0 else 'Heading 1');doc.add_paragraph(purpose)
 para=doc.add_paragraph();para.paragraph_format.space_before=Pt(8);para.paragraph_format.space_after=Pt(16)
 if crop:
  ow,oh=Image.open(img).size;x,y,w,h=crop
  shape=para.add_run().add_picture(str(img),width=Cm(17.5),height=Cm(17.5*h/w))
  src=OxmlElement('a:srcRect')
  for key,val in [('l',x/ow),('t',y/oh),('r',(ow-x-w)/ow),('b',(oh-y-h)/oh)]:src.set(key,str(round(val*100000)))
  fill=shape._inline.xpath('.//pic:blipFill')[0];fill.insert(1,src)
 else:para.add_run().add_picture(str(img),width=Cm(17.5))
 p=doc.add_paragraph();p.add_run('照圖做').bold=True
 doc.add_paragraph(action)
 p=doc.add_paragraph(note.replace('没','沒'));p.paragraph_format.space_before=Pt(7)
 for z in p.runs:z.font.size=Pt(11);z.font.color.rgb=RGBColor.from_string('526966')
 if i==1:
  p=doc.add_paragraph('官方依據：developers.cloudflare.com/fundamentals/account/create-account/','Caption')
 if i==0:doc.add_paragraph('有帳號就沿用；缺帳號，Agent 會開官方頁帶你辦。','Caption')
for el in doc.styles.element.xpath('.//w:pBdr'): el.getparent().remove(el)
for el in doc.element.xpath('.//w:pBdr'): el.getparent().remove(el)
out=R/'docs/AI秘書完整操作圖解.docx';doc.save(out)
(B/'guide-pages.json').write_text(json.dumps([{'step':i+1,'title':p[0],'image':str(p[2]),'crop':p[3]} for i,p in enumerate(pages)],ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX created:',len(pages),'steps')
